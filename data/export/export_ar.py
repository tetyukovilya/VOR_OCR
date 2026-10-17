# -*- coding: utf-8 -*-
from __future__ import annotations

import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

from app.rules.engine_ar import VorRow

FONT_NAME = "Arial"
HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
HEADER_FONT = Font(name=FONT_NAME, bold=True, color="FFFFFF", size=10)
WARN_FILL = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

VOR_HEADERS = [
    ("Раздел", 30), ("Подраздел", 34), ("Наименование работы", 45),
    ("Ед. изм.", 10), ("Объём", 12), ("Шифр документа", 16),
    ("№ листа", 10), ("Исходное проектное значение", 30),
    ("Формула расчёта", 20), ("Статус проверки", 26), ("Комментарий", 45),
]


def export_vor_draft(rows: list[VorRow], warnings: list[str], out_path: str) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "ВОР (черновик, АР)"
    ws.sheet_view.showGridLines = False
    for col_idx, (title, width) in enumerate(VOR_HEADERS, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width
        cell = ws.cell(row=1, column=col_idx, value=title)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        cell.border = BORDER
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 30

    for r_idx, row in enumerate(rows, start=2):
        values = [
            row.razdel, row.podrazdel, row.naimenovanie_raboty, row.edinitsa,
            row.obyom, row.shifr, row.nomer_lista, row.ishodnoe_znachenie,
            row.formula, row.status, row.kommentariy,
        ]
        for c_idx, val in enumerate(values, start=1):
            cell = ws.cell(row=r_idx, column=c_idx, value=val)
            cell.font = Font(name=FONT_NAME, size=10)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = BORDER
            cell.fill = WARN_FILL

    if warnings:
        ws2 = wb.create_sheet("Предупреждения")
        ws2.column_dimensions["A"].width = 120
        for i, w in enumerate(warnings, start=1):
            c = ws2.cell(row=i, column=1, value=w)
            c.font = Font(name=FONT_NAME, size=10)
            c.alignment = Alignment(wrap_text=True)

    wb.save(out_path)


if __name__ == "__main__":
    import sys
    import os
    from app.extraction.ar import process_ar_sheet, load_rules
    from app.rules.engine_ar import build_floor_vor_rows

    if len(sys.argv) > 1:
        pdf_path = sys.argv[1]
    else:
        pdf_path = input("Укажите путь к PDF-файлу листа АР: ").strip().strip('"')

    if len(sys.argv) > 2:
        out_path = sys.argv[2]
    else:
        default_dir = os.path.join(os.path.dirname(os.path.abspath(pdf_path)), "output")
        os.makedirs(default_dir, exist_ok=True)
        out_path = os.path.join(default_dir, "ВОР_черновик_АР.xlsx")

    rules = load_rules()
    result = process_ar_sheet(pdf_path)
    rows, warnings = build_floor_vor_rows(result["sheet"], result["floor_table"], rules)
    export_vor_draft(rows, warnings, out_path)
    print("Сохранено:", out_path)
    print("Строк ВОР:", len(rows))
    for w in warnings:
        print("ВНИМАНИЕ:", w)