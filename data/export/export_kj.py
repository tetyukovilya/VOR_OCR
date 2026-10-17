"""Экспорт черновика ВОР (раздел КЖ) в Excel вместе с журналом сверки."""
from __future__ import annotations

import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

from app.rules.engine_kj import VorRow

FONT_NAME = "Arial"
HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
HEADER_FONT = Font(name=FONT_NAME, bold=True, color="FFFFFF", size=10)
STATUS_COLORS = {
    "Черновой": PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid"),
    "Требует уточнения": PatternFill(start_color="FCE4E4", end_color="FCE4E4", fill_type="solid"),
    "Проверено ГИП": PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid"),
}
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

VOR_HEADERS = [
    ("Раздел", 30), ("Подраздел", 30), ("Наименование работы", 45),
    ("Ед. изм.", 10), ("Объём", 12), ("Шифр документа", 16),
    ("№ листа", 10), ("Исходное проектное значение", 32),
    ("Формула расчёта", 20), ("Статус проверки", 16), ("Комментарий", 40),
]


def export_vor_draft(rows: list[VorRow], cross_check: dict, out_path: str) -> None:
    wb = openpyxl.Workbook()

    ws = wb.active
    ws.title = "ВОР (черновик)"
    ws.sheet_view.showGridLines = False
    for col_idx, (title, width) in enumerate(VOR_HEADERS, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width
        cell = ws.cell(row=1, column=col_idx, value=title)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        cell.border = BORDER
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 30

    for r_idx, row in enumerate(rows, start=2):
        values = [
            row.razdel, row.podrazdel, row.naimenovanie_raboty, row.edinitsa,
            row.obyom, row.shifr, row.nomer_lista, row.ishodnoe_znachenie,
            row.formula, row.status, row.kommentariy,
        ]
        fill = STATUS_COLORS.get(row.status, PatternFill())
        for c_idx, val in enumerate(values, start=1):
            cell = ws.cell(row=r_idx, column=c_idx, value=val)
            cell.font = Font(name=FONT_NAME, size=10)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = BORDER
            cell.fill = fill

    # ---------- Лист сверки ----------
    ws2 = wb.create_sheet("Журнал сверки")
    ws2.column_dimensions["A"].width = 45
    ws2.column_dimensions["B"].width = 60
    ws2.sheet_view.showGridLines = False

    title = ws2.cell(row=1, column=1, value="Автоматическая сверка: ведомость расхода стали")
    title.font = Font(name=FONT_NAME, bold=True, size=12)

    check_rows = [
        ("Сумма масс по позициям спецификации, кг", cross_check["total_from_spec_kg"]),
        ("Итого по сводной ведомости 'Всего', кг", cross_check["total_from_summary_kg"]),
        ("Расхождение, кг", cross_check["difference_kg"]),
        ("В пределах допуска?", "Да" if cross_check["within_tolerance"] else "НЕТ — требуется решение ГИП"),
        ("Позиции, учтённые в сумме", ", ".join(cross_check["positions_used"])),
    ]
    r = 3
    for label, val in check_rows:
        c1 = ws2.cell(row=r, column=1, value=label)
        c1.font = Font(name=FONT_NAME, bold=True, size=10)
        c2 = ws2.cell(row=r, column=2, value=str(val))
        c2.font = Font(name=FONT_NAME, size=10)
        c2.alignment = Alignment(wrap_text=True)
        r += 1

    if not cross_check["within_tolerance"]:
        note = ws2.cell(
            row=r + 1, column=1,
            value=(
                "Вероятная причина расхождения (лист 5235-1-9-КЖ2.1): позиция 'КР1' "
                "(Каркас плоский, 231 шт., ссылка на документ 5235-1-1-КЖ2.1.И-КР1) "
                "учтена в сводной ведомости, но её масса не детализирована в "
                "'Спецификации к армированию' этого листа (только кол-во штук). "
                "Требуется решение: включать КР1 в объём армирования по этому листу "
                "или указывать отдельной позицией со ссылкой на документ КР1."
            ),
        )
        note.font = Font(name=FONT_NAME, italic=True, size=10, color="C00000")
        note.alignment = Alignment(wrap_text=True, vertical="top")
        ws2.merge_cells(start_row=r + 1, start_column=1, end_row=r + 5, end_column=2)
        ws2.row_dimensions[r + 1].height = 90

    wb.save(out_path)


if __name__ == "__main__":
    import sys
    import os
    from app.extraction.kj import process_kj_sheet
    from app.rules.engine_kj import build_vor_draft

    if len(sys.argv) > 1:
        pdf_path = sys.argv[1]
    else:
        pdf_path = input("Укажите путь к PDF-файлу листа КЖ: ").strip().strip('"')

    if len(sys.argv) > 2:
        out_path = sys.argv[2]
    else:
        default_dir = os.path.join(os.path.dirname(os.path.abspath(pdf_path)), "output")
        os.makedirs(default_dir, exist_ok=True)
        out_path = os.path.join(default_dir, "ВОР_черновик_КЖ.xlsx")

    result = process_kj_sheet(pdf_path)
    rows, check = build_vor_draft(result["sheet"], result["spec_rows"], result["steel_summary"])
    export_vor_draft(rows, check, out_path)
    print("Сохранено:", out_path)
    print("Сверка:", check)